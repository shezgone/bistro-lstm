# -*- coding: utf-8 -*-
"""대표 보고 덱 v2 (9/14) — 6장 요약본: 약자 풀이 병기, FDE 협업 중심, 기술은 도표로, 상세는 노트로"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

F = "NanumSquare"
SCALE = 1.15
INK = RGBColor(0x26,0x26,0x26); GREY = RGBColor(0x6B,0x6B,0x6B)
LINE = RGBColor(0xC9,0xC9,0xC9); NAVY = RGBColor(0x24,0x34,0x47)
GREEN = RGBColor(0x00,0x8A,0x3E); BGGREY = RGBColor(0xF5,0xF5,0xF4)
WHITE = RGBColor(0xFF,0xFF,0xFF); HL = RGBColor(0xEE,0xF5,0xF0)
WARM = RGBColor(0xB0,0x53,0x2F); BLUE = RGBColor(0x1C,0x5C,0xAB)

SC = "/private/tmp/claude-502/-Users-user-vibe-bistro-lstm/031ca4b9-003d-4cb2-afaa-06353432dc9b/scratchpad"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
PAGE = [0]
def new_slide():
    PAGE[0] += 1
    return prs.slides.add_slide(prs.slide_layouts[6])

def runs(s, x, y, w, h, lines, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp=2):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01); tf.margin_bottom = Inches(0.01)
    for i, rs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp)
        for t, b, c, z in rs:
            r = p.add_run(); r.text = t
            r.font.name = F; r.font.size = Pt(round((z if z else size) * SCALE, 1))
            r.font.bold = b; r.font.color.rgb = c if c else INK
            rp = r._r.get_or_add_rPr(); rp.append(rp.makeelement(qn('a:ea'), {'typeface': F}))
    return tf

def hline(s, x, y, w, color=LINE, pt=0.75):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x+w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(pt); return ln

def rect(s, x, y, w, h, fill=None, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False; return shp

def header(s, tag, msg):
    runs(s, 0.6, 0.32, 11.5, 0.3, [[(tag, True, GREY, 10.5)]])
    runs(s, 0.6, 0.6, 12.1, 0.6, [[(msg, True, NAVY, 17)]])
    hline(s, 0.6, 1.22, 12.13, NAVY, 1.2)

def footnote(s, text):
    runs(s, 0.6, 7.1, 11.6, 0.3, [[("주: " + text, False, GREY, 8)]])
    runs(s, 12.85, 7.1, 0.4, 0.25, [[(str(PAGE[0]), False, GREY, 9)]], align=PP_ALIGN.RIGHT)

def sidebar(s, x, y, h, color=NAVY):
    rect(s, x, y, 0.045, h, color, None)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

RULE = "채점 규약(전 장 공통): 각 시점에 실제 존재했던 데이터만 사용(실시간 빈티지) · 정답 = GDP 속보치 · 발표 19주 전부터 매주 예측 · 2018~2025년 32개 분기 · RMSE(평균제곱근오차) — 낮을수록 정확."

# ============ 1. 표지 ============
s = new_slide()
runs(s, 0.75, 0.55, 8.0, 0.35, [[("네이버클라우드 AX Forward Lab", True, GREY, 12)]])
runs(s, 0.75, 2.4, 11.9, 1.8, [
    [("한국은행 GDP 나우캐스팅 협업 — 3개월의 여정", True, NAVY, 31)],
    [("", False, INK, 8)],
    [("나우캐스팅(nowcasting) = 이번 분기 성장률을 공식 발표 전에 실시간 추정하는 일", False, GREY, 13)],
    [("GitHub 커밋으로 일한 FDE 협업의 기록 — 성과, 일하는 방식, 그리고 10월의 무대", False, INK, 16)],
])
hline(s, 0.78, 4.5, 4.2, NAVY, 1.6)
runs(s, 0.78, 4.7, 11.2, 0.8, [
    [("한국은행 기존 모형 대비 예측 오차 2.3% 개선 · 한은과 결과 합의 완료 · 10월 총재님·임원진 보고 일정 확정", False, GREY, 12)],
], sp=3)
runs(s, 0.78, 6.2, 11.9, 0.35, [[("한은–네이버 2026년 협력과제 8건 중 우선순위 2위 「시계열 Track1 (거시경제 예측)」 — 한은 담당: 이창훈·윤희근 (디지털혁신실)", False, GREY, 10.5)]])
runs(s, 0.78, 6.58, 8.0, 0.35, [[("2026. 9. 14.  |  대표 보고  |  작성: 김용민  |  FDE = Forward Deployed Engineering, 고객 현장 밀착 엔지니어링", False, GREY, 10.5)]])
notes(s, "6장 요약본(9/1 11장본의 개정판 — 상세 수치와 통계 논거는 각 장의 노트로 이동). 발표 10분 기준. 대표님은 한은 FDE 전사 사례(우선순위 1위 규정 온톨로지, 코딩 어시스턴트 등)를 이미 보고받으심 — 본 건은 그 지도에서 우선순위 2위 '시계열 Track1'의 상세 보고. 시계열 Track2(뉴스 인과관계, 데이터사이언스팀)와는 별개 트랙.")

# ============ 2. 요약 ============
s = new_slide()
header(s, "1  |  요약", "세 가지만 기억해 주시면 됩니다")
blocks = [
    ("성과", GREEN, "한국은행 기존 모형보다 예측 오차(RMSE·평균제곱근오차)를 2.3% 줄인 구성을 찾았고, 9월 3일 회의에서 한은이 이 결과에 합의했습니다. 실패한 실험까지 전부 기록으로 남긴 3개월의 결론입니다."),
    ("방식", NAVY, "모든 협업이 GitHub 저장소의 커밋으로 이루어졌습니다. 한은이 데이터와 채점 기준을 올리면 당사가 검증과 실험으로 회신했고, 서로의 결과를 서로의 코드로 재현해 확인했습니다 — FDE 방식의 실제 모습입니다."),
    ("다음", WARM, "협업의 성격이 \"모형 검증 지원\"에서 \"공동 연구\"로 넓어졌습니다. 10월 1일 한은 중간세미나(행원 대상), 10월 26일 주간현안포럼(총재님·임원진 보고)이 확정됐고, 공동 논문을 협의 중입니다."),
]
y = 1.5
for t, c, d in blocks:
    sidebar(s, 0.6, y, 0.95, c)
    runs(s, 0.85, y + 0.05, 1.5, 0.4, [[(t, True, c, 14)]])
    runs(s, 2.45, y + 0.02, 10.2, 0.95, [[(d, False, INK, 11)]])
    y += 1.18
runs(s, 0.6, 5.12, 6.0, 0.3, [[("석 달의 흐름", True, GREY, 10)]])
tl = [
    ("6월", "착수 — 한은 현행 시스템을 실시간 데이터로 재현"),
    ("7월", "국면전환 연구 — 성능 최고였으나 채택 보류"),
    ("8월 초", "한은 채점 개정 — 당일 전면 재검증 대응"),
    ("8월 말", "적응학습(LoRA) 실증 — 32분기 최고 기록"),
    ("9월", "결과 합의(3일 회의) · 본 보고"),
    ("10월", "중간세미나(1일) · 총재님 보고(26일)"),
]
x = 0.6
for when, d in tl:
    rect(s, x, 5.42, 1.85, 1.35, BGGREY, None)
    runs(s, x + 0.09, 5.5, 1.67, 0.3, [[(when, True, NAVY, 10.5)]])
    runs(s, x + 0.09, 5.82, 1.67, 0.9, [[(d, False, INK, 8.3)]])
    x += 2.05
footnote(s, RULE)
notes(s, "질문 대비 수치: 기준선 = 한은 개정판 XGBoost 0.750, 최종 구성 0.7326(-2.3%), 32개 분기 중 18개 분기 승. 통계적 수위는 4장 노트 참조. 7월 국면전환 연구(0.718, 전 여정 최고 성능)는 '실시간으로 경기 국면을 단정할 수 있는가'라는 중앙은행 설명책임 요건과 충돌해 보류 — 그 직관은 '분기 초 6주 교대'라는 수용 가능한 형태로 최종 구성에 계승됨.")

# ============ 3. 일하는 방식 (FDE) ============
s = new_slide()
header(s, "2  |  일하는 방식", "저장소 하나를 가운데 두고 일했습니다 — 데이터, 코드, 결과가 모두 커밋으로")
s.shapes.add_picture(f"{SC}/fde_collab.png", Inches(0.5), Inches(1.5), width=Inches(6.9))
X, W = 7.7, 5.05
runs(s, X, 1.55, W, 0.35, [[("실제 장면 셋", True, NAVY, 12.5)]])
scenes = [
    ("한은이 채점 코드의 결함을 스스로 개정", "당사는 당일로 40여 개 모형을 전부 다시 채점해 회신했습니다. 당사의 종전 최고 기록이 무효가 되는 것도 함께 공개했습니다."),
    ("한은이 새 실험 코드 3만 7천 줄을 커밋", "당사가 이튿날 핵심 진단을 재현하고, 사흘 만에 한은이 제안만 했던 적응학습을 첫 실증으로 답했습니다."),
    ("당사 종합리포트를 한은이 자기 저장소에 커밋", "산출물이 고객의 공식 기록으로 편입됐습니다 — 문서가 아니라 커밋으로 받은 수신 확인입니다."),
]
y = 2.05
for t, d in scenes:
    sidebar(s, X, y, 1.28, GREEN)
    runs(s, X + 0.2, y, W - 0.2, 0.32, [[(t, True, INK, 10.5)]])
    runs(s, X + 0.2, y + 0.34, W - 0.2, 0.9, [[(d, False, GREY, 9.5)]])
    y += 1.5
rect(s, 0.6, 6.55, 12.13, 0.45, BGGREY, None)
sidebar(s, 0.6, 6.55, 0.45, NAVY)
runs(s, 0.85, 6.62, 11.7, 0.35, [[("이 틀 — ① 데이터 기준(빈티지) 통일  ② 전처리 코드 공유  ③ 상호 재현 저장소 — 은 고객과 과제가 달라져도 그대로 쓸 수 있는 FDE 협업 템플릿입니다.", True, NAVY, 10.5)]])
footnote(s, "빈티지(vintage) = 각 시점에 실제로 존재했던 데이터의 스냅샷 — 나중에 개정된 값을 미리 쓰는 착시를 차단. 당사 저장소 커밋 90여 건(6~9월).")
notes(s, "정확한 구조(질문 대비): 단일 공용 저장소가 아니라 양측 저장소를 상호 구독(pull)하는 구조 — 당사는 한은 저장소에 푸시하지 않음(권한·기밀 분리). 한은 빈티지 데이터는 당사 저장소에 커밋하지 않는 규칙 유지(기밀 관리). 'FDE = 고객 코드를 읽는 것이 요구사항 분석': 8월 초 개정 커밋의 diff에서 단위 결함의 의미를, 8월 말 커밋에서 실험 설계와 한계까지 독해해 대응함.")

# ============ 4. 결과 ============
s = new_slide()
header(s, "3  |  결과", "기존 모형 대비 오차 2.3% 개선 — \"확정 우위\"가 아니라 \"개선 방향\"으로 보고했습니다")
s.shapes.add_picture(f"{SC}/ceo_final2.png", Inches(0.45), Inches(1.7), width=Inches(7.5))
rect(s, 0.6, 5.25, 7.2, 1.15, BGGREY, None)
runs(s, 0.78, 5.35, 6.9, 1.0, [
    [("\"분기 초 6주\"란 — ", True, INK, 10), ("공식 월간통계는 약 한 달 늦게 나오므로, 분기가 시작되고", False, INK, 10)],
    [("6주 동안은 그 분기의 통계가 한 건도 없습니다. 통계 기반 모형이 힘을 못 쓰는", False, INK, 10)],
    [("이 \"지표 공백\" 구간에만 파운데이션 모델을 씁니다.", False, INK, 10)],
], sp=2)
X, W = 8.15, 4.6
sidebar(s, X, 1.7, 2.15, GREEN)
runs(s, X + 0.2, 1.7, W - 0.2, 2.1, [
    [("최종 구성 (말로 풀면)", True, NAVY, 12)],
    [("평소에는 한은의 XGBoost(트리 기반 머신러닝) 모형을", False, INK, 10.5)],
    [("그대로 쓰고, 분기 초 6주만 파운데이션 모델로 교대.", False, INK, 10.5)],
    [("적응 사례가 12분기 넘게 쌓인 뒤에는 적응학습(LoRA)", False, INK, 10.5)],
    [("부품으로 업그레이드합니다.", False, INK, 10.5)],
], sp=2.5)
sidebar(s, X, 4.15, 2.3, NAVY)
runs(s, X + 0.2, 4.15, W - 0.2, 2.3, [
    [("왜 \"개선 방향\"까지만 말하나", True, NAVY, 12)],
    [("표본이 32개 분기뿐이라 이 크기의 차이는 통계적으로", False, INK, 10.5)],
    [("확정할 수 없습니다 (Diebold-Mariano 검정, 두 모형의", False, INK, 10.5)],
    [("예측력 차이 검정 — p=0.185). 한은에도 같은 수위로", False, INK, 10.5)],
    [("보고했고, 이 정직한 수위가 신뢰의 바탕이 됐습니다.", False, INK, 10.5)],
], sp=2.5)
footnote(s, RULE)
notes(s, "통계 상세(질문 대비): 검정력 분석상 이 표본(유효 32개)에서는 6.2% 개선도 p≈0.09 — 애초 '확정 우위'가 성립할 수 없는 표본 크기. 한은 자체 Model Confidence Set 검정도 비교 12개 모형 전원 잔류로 같은 결론. 최종 구성은 32분기 중 18분기 승. 차트 회색 막대(국면전환 0.718): 성능은 전 여정 최고 — (구두 대응 전용) 대표님이 과거 언급하신 HMM(은닉 마르코프 모형)과 같은 계보의 접근. 질문이 나오면: '그 방향이 성능으로는 옳았고, 실시간 국면 단정의 수용성 문제 때문에 분기 초 교대라는 형태로 번역해 계승했다'로 답변. 선제 언급은 하지 않음.")

# ============ 5. 기술 노트 ============
s = new_slide()
header(s, "4  |  기술 노트", "파운데이션 모델을 쓰는 두 가지 방법 — 즉석 사용과 가벼운 적응")
runs(s, 0.6, 1.4, 12.1, 0.55, [
    [("파운데이션 모델(foundation model) = 세계 각국의 시계열 수백만 개로 미리 학습해 둔 범용 예측 모델. 이번에는 Chronos-2(아마존)와 BISTRO(국제결제은행 BIS의 거시경제 특화 모델) 두 개를 같은 조건으로 실측했습니다.", False, INK, 10.5)],
])
s.shapes.add_picture(f"{SC}/icl_lora.png", Inches(0.45), Inches(2.1), width=Inches(6.8))
X, W = 7.5, 5.25
sidebar(s, X, 2.15, 2.0, NAVY)
runs(s, X + 0.2, 2.15, W - 0.2, 2.0, [
    [("왜 적응이 필요했나", True, NAVY, 12)],
    [("즉석 사용만으로는 새 지표가 발표돼도 예측이 잘 바뀌지", False, INK, 10.5)],
    [("않는 문제가 있었습니다 (한은이 먼저 지적한 현상).", False, INK, 10.5)],
    [("가볍게 적응시키자 이 문제가 크게 줄었고, 두 모델 모두", False, INK, 10.5)],
    [("오차가 6~9% 개선됐습니다.", False, INK, 10.5)],
], sp=2.5)
sidebar(s, X, 4.5, 2.2, WARM)
runs(s, X + 0.2, 4.5, W - 0.2, 2.2, [
    [("과적합(문제 암기)을 막는 안전장치", True, WARM, 12)],
    [("· 같은 학습을 3회 반복해 평균만 신뢰 (운 좋은 1회 배제)", False, INK, 10.5)],
    [("· 학습에 쓴 분기는 채점에서 제외", False, INK, 10.5)],
    [("· 사례가 12분기 넘게 쌓인 뒤에만 사용 — 그 전에는", False, INK, 10.5)],
    [("  오히려 나빠진다는 것을 실측으로 확인한 규칙입니다", False, INK, 10.5)],
], sp=2.5)
footnote(s, "ICL(In-Context Learning·문맥 내 학습) · LoRA(Low-Rank Adaptation·저계수 적응). 학습 데이터는 모두 \"그 시점에 알 수 있던 값\"만 사용 — 미래 정보 유입 원천 차단.")
notes(s, "수치 상세(질문 대비): LoRA 적응 후(2021~2025 평가창, seed 3회 예측 평균) BISTRO 0.596→0.560(-6.0%), Chronos-2 0.619→0.564(-8.9%). '새 지표 반영' 정량화 = 예측 초반 6주 대비 다음 6주의 오차 감소율 — BISTRO는 적응 전 역행(+10.1%)이 적응 후 +3.2%로 완화. BISTRO는 seed 간 편차가 커(±0.028) 단일 실행을 믿지 않고 3회 평균만 보고 — 이 원칙이 두 차례 성급한 결론을 걸러냄. BISTRO는 BIS Working Paper 1337의 공개 모델을 공개 체크포인트 그대로, 당사 채점 규약으로 실측한 것.")

# ============ 6. 다음 단계 ============
s = new_slide()
header(s, "5  |  다음 단계", "10월, 한국은행의 공식 무대 두 곳에 섭니다")
items = [
    ("확정 일정", GREEN, "10월 1일 중간세미나(행원 대상) · 10월 26일 주간현안포럼(총재님·임원진 보고) — 9월 3일 회의에서 결과 합의와 함께 확정된 일정입니다. 세미나 자료는 당사가 초안을 만들어 사전 공유합니다."),
    ("공동 논문", NAVY, "\"사전학습 시계열 모델은 언제 가치를 더하는가\" — 실측으로 확인한 기여 5건에 적응학습 장(章)을 더해, 한은 워킹페이퍼를 먼저 내고 국제 학술지에 투고하는 구도로 협의 중입니다."),
    ("협업 확대", NAVY, "한은이 먼저 제안한 적응학습 공동 연구의 역할 분담을 협의 중 — 당사는 Chronos-2 트랙과 적응 안정화를 맡는 방향입니다."),
    ("FDE 틀의 재사용", WARM, "이번에 검증한 협업 틀(데이터 기준 통일 → 전처리 코드 공유 → 상호 재현 저장소)은 고객과 과제가 달라져도 그대로 쓸 수 있습니다 — 전사 FDE 사례로 일반화할 수 있는 자산입니다."),
]
y = 1.5
for t, c, d in items:
    sidebar(s, 0.6, y, 1.0, c)
    runs(s, 0.85, y + 0.03, 2.3, 0.4, [[(t, True, NAVY, 12.5)]])
    runs(s, 3.25, y + 0.03, 9.4, 0.95, [[(d, False, INK, 10.5)]])
    y += 1.22
runs(s, 0.6, 6.5, 12.0, 0.4, [[("출발점이었던 \"모델 해석 질의\"에 3개월 만에 실측과 개선으로 답했고, 이제 고객의 공식 일정에 함께 섭니다.", True, NAVY, 12)]])
footnote(s, RULE)
notes(s, "(구두 전용 — 문서에는 쓰지 않음) ① 10/26 주간현안포럼의 총재님은 BISTRO 논문(BIS WP 1337)의 1저자(BIS 재직 시). 디지털혁신실의 논문 해석 질의가 본 프로젝트의 기점이었고, 원 논문 모델의 동일 규약 실측과 적응 개선으로 답한 수미상관 구조 — BISTRO 관련 서술은 항상 '동일 규약·공개 체크포인트 실측' 프레임 유지. ② 조직 지형: 디지털혁신실(협업 파트너, 딥러닝 시스템 개선)과 경제모형실(전문가 판단 전망, 최근 산학협력으로 BISTRO 연구)은 경쟁 구도로 보임 — 산학 파트너 실명 포함 문서 명시 금지. ③ 논문 저자 구성 등 회사 차원 조율이 필요해지면 별도 보고.")

out = "/Users/user/vibe/bistro-lstm/docs/한은협업_여정보고_대표_2026-09-14.pptx"
prs.save(out); print("saved:", out)
